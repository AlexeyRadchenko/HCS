from datetime import datetime, date
from openpyxl import load_workbook
from openpyxl.styles import Border, Side, Alignment, Font, NamedStyle
from jinja2 import Environment, BaseLoader
from locale import setlocale, LC_TIME
from pathlib import Path
from ..settings.settings import settings
from ..utils.utils import calcSum, sorting_works_group
from ..database.mkd_works.models import YearActfiles, MonthActfiles
from ..database.mkd_works.crud import (create_mkd_works_db_object, update_bg_task_status, get_mainwork_by_id,
    get_subwork_by_id, get_fixwork_by_id)

MONTHS_MAP = {
    1: 'январь',
    2: 'февраль',
    3: 'март',
    4: 'апрель',
    5: 'май',
    6: 'июнь',
    7: 'июль',
    8: 'август',
    9: 'сентябрь',
    10: 'октябрь',
    11: 'ноябрь',
    12: 'декабрь'
}


def get_double_index(doubles_list, search_id):
    for index, id in enumerate(doubles_list):
        if id == search_id:
            return index


def apply_style(cell, style):
    """Применяет стиль к ячейке."""
    cell.style = style


def create_group(main_work, work):
    """Создает новую группу для работы."""
    work_id = work.get('mainwork_id') if work.get('mainwork_id') else work.get('subwork_id') if work.get('subwork_id') else work.get('fixwork_id')
    return {
        'main_work': main_work.work[3:],
        'main_work_id': main_work.id,
        'total_sum': work['sum'],
        'works': [{
            'work_id': work_id,
            'name_work': work['name_work'],
            'period': work['act_custom_period'],
            'quantity': work['quantity'],
            'unitcost': work['unitcost'],
            'sum': work['sum'],
        }],
        'record_ids': [work_id],
    }

def find_group_index(groups, main_work_id):
    """Находит индекс группы по ID главной работы."""
    for index, group in enumerate(groups):
        if group['main_work_id'] == main_work_id:
            return index
    return None


async def grouping_and_sum_works(works, db_session):
    """Группирует и суммирует работы."""
    group_works = []
    group = None
    groups_exist_ids = []

    for work in works:
        main_work = await get_mainwork_by_id(db_session, work['mainwork_id'])
  
        if main_work.id not in groups_exist_ids:
            groups_exist_ids.append(main_work.id)
            group = create_group(main_work, work)
            group_works.append(group)
            #print("TOTAL SUM ON CREATE", group['total_sum'])
        else:
            index = find_group_index(group_works, main_work.id)
            if index is None:
                raise ValueError(f"Группа с main_work_id={main_work.id} не найдена")
            group = group_works[index]
            #print(">>>>>>>>>>>>>>>>>>", group['total_sum'], work['sum'])
            group['total_sum'] = calcSum(*[work['sum']], sum=group['total_sum'])
            #print("<<<<<<<<<<<<<<<<<<<<<<<<", group['total_sum'])
            work_id = work.get('mainwork_id') if work.get('mainwork_id') else work.get('subwork_id') if work.get('subwork_id') else work.get('fixwork_id')
            group['works'].append({
                'work_id': work_id,
                'name_work': work['name_work'],
                'period': work['act_custom_period'],
                'quantity': work['quantity'],
                'unitcost': work['unitcost'],
                'sum': work['sum'],
            })
            group['record_ids'].append(work_id)
            #print("TOTAL SUM AFTER CREATE", group['total_sum'])
            
            """if work['id'] in group['record_ids']:
                # Если работа уже существует в группе, обновляем только её данные
                double_index = get_double_index(group['record_ids'], work['id'])
                group['works'][double_index]['quantity'] += work['quantity']
                print("!@@@@@@@@", group['works'][double_index]['sum'])
                print(work['sum'])
                print(calcSum(*[group['works'][double_index]['sum'], work['sum']]))
                group['works'][double_index]['sum'] = calcSum(*[group['works'][double_index]['sum'], work['sum']])
            else:
                # Если работа новая, добавляем её в группу и обновляем total_sum
                group['total_sum'] += float(work['sum'])
                group['works'].append({
                    'work_id': work['id'],
                    'name_work': work['work'],
                    'period': work['period'],
                    'quantity': work['quantity'],
                    'unitcost': work['unitcost'],
                    'sum': work['sum'],
                })
                group['record_ids'].append(work['id'])"""

   
            
    #print("BLYAAAAAAAAAAAAAAAAA", [(g['main_work_id'], len(g['works'])) for g in group_works])    
    return group_works

async def group_by_company_works_type(works):
    groups = []
    existing_groups = []
    for work in works:
        if work['work_company_type'] not in existing_groups:
            existing_groups.append(work['work_company_type'])
            groups.append({
                'group': work['work_company_type'],
                'works': [work],
                'group_sum': work['sum'],
            })
        else:
            index = existing_groups.index(work['work_company_type'])
            if index is None:
                raise ValueError(f"Группа с work_company_type={work['work_company_type']} не найдена")
            groups[index]['works'].append(work)
            groups[index]['group_sum'] = calcSum(*[work['sum']], sum=groups[index]['group_sum'])
    return groups

def write_header_data(doc_sheet, write_data, jinja_env):
    """Записывает заголовок в документ."""
    end_row_num = 1
    for idx, xlsx_row in enumerate(doc_sheet.iter_rows(), start=1):
        if idx < 24:
            end_row_num = idx
            for xlsx_cell in xlsx_row:
                if xlsx_cell and isinstance(xlsx_cell.value, str):
                    j_template = jinja_env.from_string(xlsx_cell.value)
                    xlsx_cell.value = j_template.render(write_data)
    return end_row_num


def write_header_data_month_act(doc_sheet, write_data, jinja_env):
    """Записывает заголовок в документ."""
    for idx, xlsx_row in enumerate(doc_sheet.iter_rows(), start=1):
        if idx == 1:
            j_template = jinja_env.from_string(xlsx_row[1].value)
            xlsx_row[1].value = j_template.render(write_data)
        if idx == 2:
            j_template = jinja_env.from_string(xlsx_row[2].value)
            xlsx_row[2].value = j_template.render(write_data)    


async def write_table_data(doc_sheet, write_data, start_write_row_num, db_session, year):
    """Записывает данные таблицы в документ."""
    # Настраиваем стили
    thin = Side(border_style="thin", color="000000")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    border_left = Border(top=thin, left=thin, right=thin, bottom=None)
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alignment_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    font = Font(name='Times New Roman', size=12)
    font_bold = Font(name='Times New Roman', size=12, bold=True)

    cell_style = NamedStyle(name="styled_cell", border=border, alignment=alignment, font=font)
    cell_style_font_bold_only = NamedStyle(name="styled_cell_font_bold_only", font=font_bold)
    cell_style_font_only = NamedStyle(name="styled_cell_font_only", font=font)
    cell_style_left = NamedStyle(name="styled_cell_left", border=border_left, alignment=alignment_left, font=font_bold)
    cell_styel_work_group = NamedStyle(name="styled_cell_work_group", alignment=alignment, font=font_bold)

    write_row = start_write_row_num + 3
    #print("WRITE DATA LENGTH ", len(write_data))
    group_works_for_sort = await grouping_and_sum_works(write_data, db_session)
    group_works = sorting_works_group(group_works_for_sort)
    #print("group LENGTH ", len(group_works))
    total_sum_of_groups = 0.00
    for group in group_works:
        doc_sheet[f'A{write_row}'].value = group['main_work']
        doc_sheet.merge_cells(f'A{write_row}:E{write_row}')
        apply_style(doc_sheet[f'A{write_row}'], cell_styel_work_group)
        write_row += 1
        for work in group['works']:
            doc_sheet[f'A{write_row}'].value = work['name_work']
            doc_sheet[f'B{write_row}'].value = work['period']
            doc_sheet[f'C{write_row}'].value = work['quantity']
            doc_sheet[f'D{write_row}'].value = work['unitcost']
            doc_sheet[f'E{write_row}'].value = work['sum']
            for col in ['A', 'B', 'C', 'D', 'E']:
                apply_style(doc_sheet[f'{col}{write_row}'], cell_style)
            write_row += 1
        #print("AASSSSSSSSSSSSSSSSSSSS", group)
        doc_sheet[f'A{write_row}'].value = 'Итого:'
        doc_sheet.merge_cells(f'A{write_row}:D{write_row}')
        #print("SUMMMMMMMMMMMMMMMMMMMMMM", group['total_sum'])
        doc_sheet[f'E{write_row}'].value = group['total_sum']
        apply_style(doc_sheet[f'E{write_row}'], cell_style_font_only)
        apply_style(doc_sheet[f'A{write_row}'], cell_style_left)
        total_sum_of_groups = calcSum(*[group['total_sum']], sum=total_sum_of_groups)
        write_row += 1

    doc_sheet[f'A{write_row}'].value = f'Итого расходы за {year}:'
    doc_sheet.merge_cells(f'A{write_row}:D{write_row}')
    doc_sheet[f'E{write_row}'].value = total_sum_of_groups
    apply_style(doc_sheet[f'E{write_row}'], cell_style_font_bold_only)
    apply_style(doc_sheet[f'A{write_row}'], cell_style_left)
    return write_row

async def write_table_data_month_act(doc_sheet, write_data, start_write_row_num, month):
    thin = Side(border_style="thin", color="000000")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    border_left = Border(top=thin, left=thin, right=thin, bottom=None)
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alignment_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    alignment_right = Alignment(horizontal="right", vertical="center", wrap_text=True)
    font = Font(name='Times New Roman', size=12)
    font_bold = Font(name='Times New Roman', size=12, bold=True)
    font_italic = Font(name='Times New Roman', size=12, italic=True, bold=True)

    cell_style = NamedStyle(name="styled_cell", border=border, alignment=alignment, font=font)
    cell_style_font_bold_only = NamedStyle(name="styled_cell_font_bold_only", font=font_bold, alignment=alignment_right)
    cell_style_font_only = NamedStyle(name="styled_cell_font_only", font=font)
    cell_style_left = NamedStyle(name="styled_cell_left", border=border_left, alignment=alignment_left, font=font_bold)
    cell_styel_work_group = NamedStyle(name="styled_cell_work_group", alignment=alignment, font=font_bold)
    cell_style_font_bold_italic = NamedStyle(name="styled_cell_font_bold_italic", font=font_italic)

    groups = await group_by_company_works_type(write_data)
    for group in groups:
        group['works'] = sorted(group['works'], key=lambda x: x['house'])

    num_work_order = 1
    total_sum_of_groups = 0.00
    for group in groups:
        doc_sheet[f'A{start_write_row_num}'].value = group['group']
        doc_sheet.merge_cells(f'A{start_write_row_num}:E{start_write_row_num}')
        apply_style(doc_sheet[f'A{start_write_row_num}'], cell_styel_work_group)
        start_write_row_num += 1
        for work in group['works']:
            #print("START ROW NUM: ", start_write_row_num)
            doc_sheet[f'A{start_write_row_num}'].value = num_work_order
            doc_sheet[f'B{start_write_row_num}'].value = work['house']
            doc_sheet[f'C{start_write_row_num}'].value = work['name_work']
            doc_sheet[f'D{start_write_row_num}'].value = work['smeta_num']
            doc_sheet[f'E{start_write_row_num}'].value = work['sum']
            for col in ['A', 'B', 'C', 'D', 'E']:
                apply_style(doc_sheet[f'{col}{start_write_row_num}'], cell_style)
            start_write_row_num += 1
            num_work_order += 1
        doc_sheet[f'A{start_write_row_num}'].value = f'Итого: {group["group"]}'
        apply_style(doc_sheet[f'A{start_write_row_num}'], cell_style_font_bold_italic)
        doc_sheet.merge_cells(f'A{start_write_row_num}:D{start_write_row_num}')
        doc_sheet[f'E{start_write_row_num}'].value = group['group_sum']
        apply_style(doc_sheet[f'E{start_write_row_num}'], cell_style_font_bold_only)
        total_sum_of_groups = calcSum(*[group['group_sum']], sum=total_sum_of_groups)
        start_write_row_num += 1
    doc_sheet[f'A{start_write_row_num}'].value = f'Итого за {month}:'
    apply_style(doc_sheet[f'A{start_write_row_num}'], cell_style_font_bold_italic)
    doc_sheet[f'E{start_write_row_num}'].value = total_sum_of_groups
    apply_style(doc_sheet[f'E{start_write_row_num}'], cell_style_font_bold_only)    
    return start_write_row_num + 1    


def write_footer_data(doc_sheet, text_list, data_to_write, jinja_env, start_row):
    """Записывает нижний колонтитул в документ."""
    font = Font(name='Times New Roman', size=12)
    row = start_row + 2

    for idx, text in enumerate(text_list, start=1):
        j_template = jinja_env.from_string(text)
        value = j_template.render(data_to_write)
        doc_sheet[f'A{row}'].value = value
        doc_sheet[f'A{row}'].font = font
        doc_sheet.merge_cells(f'A{row}:E{row}')
        row += 2 if idx != 4 else 5

def write_footer_data_month_act(doc_sheet, text_list, data_to_write, jinja_env, start_row):
    """Записывает нижний колонтитул в документ."""
    font = Font(name='Times New Roman', size=12)
    row = start_row + 2

    for idx, text in enumerate(text_list, start=1):
        j_template = jinja_env.from_string(text)
        value = j_template.render(data_to_write)
        doc_sheet[f'A{row}'].value = value
        doc_sheet[f'A{row}'].font = font
        doc_sheet.merge_cells(f'A{row}:E{row}')

async def genereate_year_act_xlsx_file_v2(year, house, data, task_uuid, db_session):
    """Генерирует годовой акт в формате XLSX."""
    setlocale(LC_TIME, 'ru_RU.UTF-8')
    env = Environment(loader=BaseLoader, autoescape=False)

    company_director = '{0} {1}{2}'.format(
        data[0]['houses']['companies']['dirsurname'], data[0]['houses']['companies']['dirname'], data[0]['houses']['companies']['dirsecondname']
    )
    company_name = data[0]['houses']['companies']['full_name']
    all_works_lst = []
    for act in data:
        if act['mainworks_details'] != []:
            for index, mainwork in enumerate(act['mainworks_details']):
                mainwork['mainwork_id'] = act['mainworks'][index]['id']
                mainwork['name_work'] = act['mainworks'][index]['work']
            all_works_lst.extend(act['mainworks_details'])        
        if act['subworks_details'] != []:
            for index, subwork in enumerate(act['subworks_details']):
                subwork['mainwork_id'] = act['subworks'][index]['mainwork_id']
                subwork['name_work'] = act['subworks'][index]['work']
            all_works_lst.extend(act['subworks_details'])
        if act['fixworks_details'] != []:
            for index, fixwork in enumerate(act['fixworks_details']):
                fixwork['mainwork_id'] = act['fixworks'][index]['mainwork_id']
                fixwork['name_work'] = act['fixworks'][index]['work']
            all_works_lst.extend(act['fixworks_details'])

    #print("LENGTH DATA@@@@@@@@@@@@@@@@@@@@@@@", len(all_works_lst))
    """for d in all_works_lst:
        print("DATA@@@@@@@@@@@@@@@@@@@@@@@", d)"""
    template_xlsx = load_workbook(settings.YEAR_ACT_FILE_TEMPLATE_PATH + settings.YEAR_ACT_FILE_TEMPLATE_NAME)
    template_sheet = template_xlsx.active
    now = datetime.now()

    template_data = {
        'act_num': f'1-{year.year}',
        'street': data[0]['houses']['street'],
        'street_number': data[0]['houses']['number'],
        'act_date': now.strftime("%-d %B %Y г."),
        'house_director': data[0]['director'] or 'Совет МКД не создан',
        'house_director_appartment': data[0]['director_appartment'] or '--',
        'company_director': company_director,
        'company_name': company_name,
    }

    footer_text = [
        '2. Всего за период с {{year_start_date}} года по {{year_end_date}} года выполнено работ (оказано услуг) на общую сумму {{total_sum}} рубл.',
        '3. Работы (услуги) выполнены (оказаны) полностью, в установленные сроки, с надлежащим качеством.',
        '4. Претензий по выполнению условий Договора Стороны друг к другу не имеют.',
        'Настоящий акт составлен в 2-х экземплярах, имеющих одинаковую юридическую силу, по одному для каждой из сторон.',
        'Подписи Сторон:',
        'Исполнитель:',
        'Генеральный директор________________________ {{company_director}}',
        'Заказчик:',
        'Председатель Совета дома___________________________ {{house_director}}'
    ]

    footer_template_data = {
        'company_director': company_director,
        'house_director': data[0]['director'] if data[0]['director'] else 'Совет МКД не создан',
        'year_start_date': date(year.year, 1, 1).strftime("%-d %B %Y"),
        'year_end_date': date(year.year, 12, 31).strftime("%-d %B %Y"),
        'total_sum': "{0:.2f}".format(float(calcSum(*[float(work['sum']) for work in all_works_lst]))),
    }

    end_write_row_num = write_header_data(template_sheet, template_data, env)
    end_write_table_row_num = await write_table_data(template_sheet, all_works_lst, end_write_row_num, db_session, year.year)
    write_footer_data(template_sheet, footer_text, footer_template_data, env, end_write_table_row_num)

    ready_file_path = f'{settings.YEAR_ACT_FILE_READY_PATH}/{house}/{year.year}/'
    ready_file_name = f'{task_uuid}_{year.year}_годовой_акт_вр_{data[0]["houses"]["street"]}-{data[0]["houses"]["number"]}.xlsx'
    db_file_name = f'{year.year}_годовой_акт_выполненных_работ_{data[0]["houses"]["street"]}-{data[0]["houses"]["number"]}.xlsx'

    Path(ready_file_path).mkdir(parents=True, exist_ok=True)
    template_xlsx.save(ready_file_path + ready_file_name)

    year_act_obj = YearActfiles(
        name=db_file_name,
        num=str(1),
        date=now,
        year=year,
        extention='xlsx',
        url='',
        path=ready_file_path + ready_file_name,
        size='',
        filetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        house_id=data[0]['houses']['id'],
    )
    await create_mkd_works_db_object(db_session, year_act_obj)
    task_ready_time = datetime.now()
    await update_bg_task_status(db_session, task_uuid, 'done', task_ready_time)

async def genereate_month_act_xlsx_file_v2(month_year, data, task_uuid, db_session, house=None):
    """Генерирует месячный акт в формате XLSX."""
    setlocale(LC_TIME, 'ru_RU.UTF-8')
    env = Environment(loader=BaseLoader, autoescape=False)

    company_name = data[0]['houses']['companies']['full_name']
    #print(data[0], company_name)
    all_works_lst = []
    for act in data:
        if act['mainworks_details'] != []:
            for index, mainwork in enumerate(act['mainworks_details']):
                mainwork['mainwork_id'] = act['mainworks'][index]['id']
                mainwork['name_work'] = act['mainworks'][index]['work']
                mainwork['work_company_type'] = act['mainworks'][index]['companyWorkType']
                mainwork['house'] = act['houses']['street'] + '-' + act['houses']['number']
                mainwork['smeta_num'] = act['num']
            all_works_lst.extend(act['mainworks_details'])        
        if act['subworks_details'] != []:
            for index, subwork in enumerate(act['subworks_details']):
                subwork['mainwork_id'] = act['subworks'][index]['mainwork_id']
                subwork['name_work'] = act['subworks'][index]['work']
                subwork['work_company_type'] = act['subworks'][index]['companyWorkType']
                subwork['house'] = act['houses']['street'] + '-' + act['houses']['number']
                subwork['smeta_num'] = act['num']
            all_works_lst.extend(act['subworks_details'])
        if act['fixworks_details'] != []:
            for index, fixwork in enumerate(act['fixworks_details']):
                fixwork['mainwork_id'] = act['fixworks'][index]['mainwork_id']
                fixwork['name_work'] = act['fixworks'][index]['work']
                fixwork['work_company_type'] = act['fixworks'][index]['companyWorkType']
                fixwork['house'] = act['houses']['street'] + '-' + act['houses']['number']
                fixwork['smeta_num'] = act['num']
            all_works_lst.extend(act['fixworks_details'])

    #print("LENGTH DATA@@@@@@@@@@@@@@@@@@@@@@@", len(all_works_lst))
    """for d in all_works_lst:
        print("DATA@@@@@@@@@@@@@@@@@@@@@@@", d)"""
    template_xlsx = load_workbook(settings.MONTH_ACT_FILE_TEMPLATE_PATH + settings.MONTH_ACT_FILE_TEMPLATE_NAME)
    template_sheet = template_xlsx.active
    #print("MONTH YEAR", month_year.strftime('%B'))
    month = MONTHS_MAP[month_year.month]
    template_header_str1 = {
        'month': MONTHS_MAP[month_year.month],
        'year': month_year.year,
        'company_name': company_name,
    }


    footer_text = [
        'инженер ПТО ________________________ '
        ]


    write_header_data_month_act(template_sheet, template_header_str1, env)
    end_write_row_num = 6
    end_write_table_row_num = await write_table_data_month_act(template_sheet, all_works_lst, end_write_row_num, month)
    write_footer_data_month_act(template_sheet, footer_text, template_header_str1, env, end_write_table_row_num)
    if house:
        ready_file_path = f'{settings.MONTH_ACT_FILE_READY_PATH}/{house}/{month_year.month}_{month_year.year}/'
    else:
        ready_file_path = f'{settings.MONTH_ACT_FILE_READY_PATH}/{month_year.year}/{month_year.month}_{month_year.year}/'
    ready_file_name = f'{task_uuid}_{month_year.month}_{month_year.year}_акт_за_{month_year.month}_{data[0]["houses"]["street"]}-{data[0]["houses"]["number"]}.xlsx'

    if house:
        house_address = f'{data[0]["houses"]["street"]}-{data[0]["houses"]["number"]}'
        db_file_name = f'{month_year.month}_{month_year.year}_акт_выполненных_работ_за_месяц_{house_address}.xlsx'
    else:
        db_file_name = f'{month_year.month}_{month_year.year}_акт_выполненных_работ_за_месяц.xlsx'

    Path(ready_file_path).mkdir(parents=True, exist_ok=True)
    template_xlsx.save(ready_file_path + ready_file_name)

    month_act_obj = MonthActfiles(
        name=db_file_name,
        num=f'{month_year.month}_{month_year.year}',
        date=datetime.now(),
        month_year=month_year,
        extention='xlsx',
        url='',
        path=ready_file_path + ready_file_name,
        size='',
        filetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        house_id=house if house else None
    )
    await create_mkd_works_db_object(db_session, month_act_obj)

    task_ready_time = datetime.now()
    await update_bg_task_status(db_session, task_uuid, 'done', task_ready_time)