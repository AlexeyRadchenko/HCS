from datetime import datetime, date
from openpyxl import load_workbook
from openpyxl.styles import Border, Side, Alignment, Font, NamedStyle
from jinja2 import Environment, BaseLoader, select_autoescape
from locale import setlocale, LC_TIME
from ..settings.settings import settings
from ..utils.utils import calcSum
from ..database.mkd_works.models import YearActfiles
from ..database.mkd_works.crud import create_mkd_works_db_object, update_bg_task_status
from pathlib import Path

def write_header_data(doc_sheet, write_data, jinja_env):
    end_row_num = 1
    for idx, xlsx_row in enumerate(doc_sheet.iter_rows(), start=1):
        if idx < 11:
            end_row_num = idx
            for xlsx_cell in xlsx_row:
                if xlsx_cell and isinstance(xlsx_cell.value, str):
                    #print("111111111111111111111", xlsx_cell, type(data[0]['houses']), data[0]['houses'])
                    j_template = jinja_env.from_string(xlsx_cell.value)
                    xlsx_cell.value = j_template.render(write_data)
                
    return end_row_num

def write_table_data(doc_sheet, write_data, start_write_row_num):
    # Настраиваем стиль границ (тонкие линии)
    thin = Side(border_style="thin", color="000000")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    # Настраиваем выравнивание текста
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Настраиваем шрифт
    font = Font(name='Times New Roman', size=12)
    # Создаем именованный стиль
    cell_style = NamedStyle(name="styled_cell")
    cell_style.border = border
    cell_style.alignment = alignment
    cell_style.font = font

    #doc_sheet.add_format(cell_style)

    write_row = start_write_row_num + 3
    for work in write_data:
        doc_sheet.cell(write_row, 1).value = work['numsprav']
        doc_sheet.cell(write_row, 1).style = cell_style
        doc_sheet.cell(write_row, 2).value = work['work']
        doc_sheet.cell(write_row, 2).style = cell_style
        doc_sheet.cell(write_row, 3).value = work['period']
        doc_sheet.cell(write_row, 3).style = cell_style
        doc_sheet.cell(write_row, 4).value = 'м2'
        doc_sheet.cell(write_row, 4).style = cell_style
        doc_sheet.cell(write_row, 5).value = work['unitcost']
        doc_sheet.cell(write_row, 5).style = cell_style
        doc_sheet.cell(write_row, 6).value = work['sum']
        doc_sheet.cell(write_row, 6).style = cell_style
        write_row += 1
    return write_row

def write_footer_data(doc_sheet, text_list, data_to_write, jinja_env, start_row):
    font = Font(name='Times New Roman', size=12)
    row = start_row + 2
    row_width = (
        doc_sheet.column_dimensions['A'].width +
        doc_sheet.column_dimensions['B'].width +
        doc_sheet.column_dimensions['C'].width + 
        doc_sheet.column_dimensions['D'].width + 
        doc_sheet.column_dimensions['E'].width +
        doc_sheet.column_dimensions['F'].width
    )
    for idx, text in enumerate(text_list, start=1):
        j_template = jinja_env.from_string(text)
        value = j_template.render(data_to_write)
        doc_sheet.cell(row, 1).value = value
        doc_sheet.width = row_width
        doc_sheet.cell(row, 1).font = font
        if idx != 4:
            row +=2
        else:
            row +=5
async def genereate_year_act_xlsx_file(year, house, data, task_uuid, db_session):
    setlocale(LC_TIME, 'ru_RU.UTF-8')
    env = Environment(
        loader=BaseLoader,
        autoescape=False,
    )
    company_director_who = '{0} {1}{2}'.format(
        data[0]['houses']['companies']['dirsurname_who_what'], data[0]['houses']['companies']['dirname'], data[0]['houses']['companies']['dirsecondname'])
    company_director = '{0} {1}{2}'.format(
        data[0]['houses']['companies']['dirsurname'], data[0]['houses']['companies']['dirname'], data[0]['houses']['companies']['dirsecondname'])
    company_name = data[0]['houses']['companies']['full_name']
    all_works_lst = []
    for act in data:
        all_works_lst+=act['mainworks']
        all_works_lst+=act['subworks']
        all_works_lst+=act['fixworks']
    template_xlsx = load_workbook(settings.YEAR_ACT_FILE_TEMPLATE_PATH + settings.YEAR_ACT_FILE_TEMPLATE_NAME)
    template_sheet = template_xlsx.active
    now = datetime.now()
    template_data = {
        'act_num': f'1-{year.year}',
        'street': data[0]['houses']['street'],
        'street_number': data[0]['houses']['number'],
        'act_date': now.strftime("%-d %B %Y г."),
        'house_director': data[0]['director'],
        'house_director_appartment': data[0]['director_appartment'],
        'company_director': company_director_who,
        'company_name': company_name,
    }
    print("!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!", template_data)
    footer_text = [
        '2. Всего за период с {{year_start_date}} года по {{year_end_date}} года выполнено работ (оказано услуг) на общую сумму {{total_sum}} рубл.',
        '3. Работы ( услуги) выполнены ( оказаны) полностью, в установленные сроки, с надлежащим качеством.',
        '4. Претензий по выполнению условий Договора Стороны друг к другу не имеют.',
        'Настоящий акт составлен в 2-х экземплярах, имеющих одинаковую юридическую силу, по одному для каждой из сторон.',
        'Подписи Сторон:',
        'Исполнитель:',
        'Генеральный директор________________________ {{company_director}}',
        'Заказчик:',
        'Председатель Совета дома___________________________ {{house_director}}'
        ]
    footer_template_data = {
        'company_director': company_director,
        'house_director': data[0]['director'],
        'year_start_date': date(year.year, 1, 1).strftime("%-d %B %Y"),
        'year_end_date': date(year.year, 12, 31).strftime("%-d %B %Y"),
        'total_sum': calcSum(*[work['sum'] for work in all_works_lst])
    }

    end_write_row_num = write_header_data(template_sheet, template_data, env)
    end_write_table_row_num = write_table_data(template_sheet, all_works_lst, end_write_row_num)
    write_footer_data(template_sheet, footer_text, footer_template_data, env, end_write_table_row_num)

                
    print(year, house, data, data[0]['houses']['street'], data[0]['houses']['number'])
    # act_num, street_name, house_number
    ready_file_path = f'{settings.YEAR_ACT_FILE_READY_PATH}/{house}/{year.year}/'
    ready_file_name = f'{task_uuid}_{year.year}_годовой_акт_вр_{data[0]['houses']['street']}-{data[0]['houses']['number']}.xlsx'
    db_file_name = f'{year.year}_годовой_акт_выполненных_работ_{data[0]['houses']['street']}-{data[0]['houses']['number']}.xlsx'
    
    Path(ready_file_path).mkdir(parents=True, exist_ok=True)
    template_xlsx.save(ready_file_path + ready_file_name)
    year_act_obj = YearActfiles(
        name=db_file_name,
        num=str(1),
        date=now,
        extention='xlsx',
        url='',
        path=ready_file_path + ready_file_name,
        size='',
        filetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        house_id=data[0]['houses']['id'],
    )
    await create_mkd_works_db_object(db_session, year_act_obj)
    task_ready_time = datetime.now()
    await update_bg_task_status(db_session, task_uuid, 'done', task_ready_time)

    

    
