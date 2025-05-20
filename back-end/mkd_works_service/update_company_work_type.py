import openpyxl
import asyncio
from datetime import datetime
from api.utils.utils import update_company_work_type


def get_work_type_id(work_type):
    id = work_type.split(",")[0]
    return int(id)

def process_xlsx(file_path):
    # Открываем xlsx файл
    wb = openpyxl.load_workbook(file_path)
    sheet_names = wb.sheetnames
    active_sheet = wb.active
    # Список для хранения всех словарей
    all_data = []


    for row in active_sheet.iter_rows(min_row=1, max_row=active_sheet.max_row, values_only=True):
        # Создаём словарь для каждой строки данных
        data = {
            'id': get_work_type_id(row[0]),
            'companyWorkType': row[1],
            'NewWorkName': row[2] if len(row) == 3 else None,
        }
        #print("DATA@@@@@@@@@@@@@@@@@@@@@@@", data)
        all_data.append(data)

    return all_data

async def main():
    file_path_main = "/home/bers/HSCUploads/mks_works/main_works.xlsx" #"house_works_accum.xlsx"  # Укажите путь к вашему файлу
    file_path_sub = "/home/bers/HSCUploads/mks_works/subworks.xlsx"
    file_path_fix = "/home/bers/HSCUploads/mks_works/main_works.xlsx"
    data_main = process_xlsx(file_path_main)
    data_sub = process_xlsx(file_path_sub)
    data_fix = process_xlsx(file_path_fix)
    #print("DATA@@@@@@@@@@@@@@@@@@@@@@@", len(data), data)
    await update_company_work_type(data_main, 'mainwork')
    await update_company_work_type(data_sub, 'subwork')
    await update_company_work_type(data_fix, 'fixwork')


if __name__ == "__main__":
    asyncio.run(main())